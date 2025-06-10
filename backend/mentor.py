from PyQt6 import QtWidgets, uic
import sys
import os
from openpyxl import load_workbook
from PyQt6.QtWidgets import QMainWindow, QMessageBox 
from backend.find_excel import get_excel_files_from_latest_vit


class MentorPage(QMainWindow):
    def __init__(self):
        super().__init__()

        # UI ve Excel dosyaları için TEK VE TUTARLI BİR app_root_path tanımlaması yapıyoruz
        if getattr(sys, 'frozen', False):
            # PyInstaller ile paketlenmişse, _MEIPASS geçici dizinin yolunu verir
            app_root_path = sys._MEIPASS
        else:
            # Normal Python betiği olarak çalışıyorsa
            # mentor.py dosyası 'backend' klasöründe.
            # Projenin ana kök dizinine (CRM_V2) ulaşmak için 'backend'den bir üst dizine çıkmalıyız.
            current_script_dir = os.path.dirname(os.path.abspath(__file__))
            app_root_path = os.path.abspath(os.path.join(current_script_dir, os.pardir))

        # UI dosyasının tam yolunu oluştur
        ui_file_name = "mentor_page_python.ui"
        ui_file_path = os.path.join(app_root_path, "ui", ui_file_name)

        # UI dosyasını yükle
        try:
            uic.loadUi(ui_file_path, self)
        except Exception as e:
            # UI yükleme hatasında kullanıcıya bilgi ver ve uygulamayı kapat
            QMessageBox.critical(self, "UI Yükleme Hatası", f"UI dosyası yüklenemedi: {ui_file_path}\nHata: {e}")
            sys.exit(1) # Uygulamayı kapat

        # Excel dosyasının yolu
        try:
            # get_excel_files_from_latest_vit fonksiyonuna doğru kök yolu (app_root_path) iletin
            latest_files = get_excel_files_from_latest_vit(app_root_path) # <-- app_root_path'i buraya verdik!

            self.excel_path = latest_files.get("Mentor")

            if not self.excel_path or not os.path.exists(self.excel_path):
                raise FileNotFoundError(f"Mentor.xlsx dosyası bulunamadı: {self.excel_path}")

        except Exception as e:
            QtWidgets.QMessageBox.critical(self, "Hata", f"Excel dosyası yüklenemedi: {str(e)}")
            self.excel_path = None # Hata durumunda excel_path'i temizle veya None yap

        # İlk açılışta multiple_tabs_changed'in tetiklenmesini engellemek için flag
        self.initializing = True

        # ComboBox'u doldur (sinyal bağlanmadan önce)
        self.fill_combobox_categories()

        # Sinyalleri bağla
        self.pushButton_ara = self.findChild(QtWidgets.QPushButton, "pushButton_ara") # UI'dan bul
        self.pushButton_tumGorusmeler = self.findChild(QtWidgets.QPushButton, "pushButton_tumGorusmeler") # UI'dan bul
        self.pushButton_tercihler = self.findChild(QtWidgets.QPushButton, "pushButton_tercihler") # UI'dan bul
        self.pushButton_exit = self.findChild(QtWidgets.QPushButton, "pushButton_exit") # UI'dan bul
        self.comboBox_cokluSekme = self.findChild(QtWidgets.QComboBox, "comboBox_cokluSekme") # UI'dan bul
        self.lineEdit_ara = self.findChild(QtWidgets.QLineEdit, "lineEdit_ara") # UI'dan bul
        self.tableWidget = self.findChild(QtWidgets.QTableWidget, "tableWidget") # UI'dan bul


        self.pushButton_ara.clicked.connect(self.search)
        self.pushButton_tumGorusmeler.clicked.connect(self.interviews)
        self.pushButton_tercihler.clicked.connect(self.go_back)
        self.pushButton_exit.clicked.connect(self.exitp)
        if self.comboBox_cokluSekme: # ComboBox'ın varlığını kontrol edin
            self.comboBox_cokluSekme.currentIndexChanged.connect(self.multiple_tabs_changed)
        
        # TableWidget'ı başlangıçta temizle (eğer UI'da zaten tanımlıysa)
        if self.tableWidget:
            self.tableWidget.setColumnCount(0)
            self.tableWidget.setRowCount(0)

        # Açılış bitti, flag kaldır
        self.initializing = False
        # Önceki sayfa referansı (geri dönmek için)
        self.sayfa = None
      
    def load_excel(self):
        if not self.excel_path or not os.path.exists(self.excel_path):
            QtWidgets.QMessageBox.warning(self, "Hata", "Mentor.xlsx dosyası bulunamadı veya yolu geçersiz.")
            return [], []

        try:
            wb = load_workbook(self.excel_path)
            sheet = wb.active
            data = []
            headers = []
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i == 0:
                    headers = list(row)
                else:
                    data.append(list(row))
            return headers, data
        except Exception as e:
            QtWidgets.QMessageBox.warning(self, "Hata", f"Excel dosyası okunamadı: {str(e)}")
            return [], []

    # ... (geri kalan fonksiyonlarınızda değişiklik yok) ...

    def display_data(self, headers, data):
        self.tableWidget.clear()
        self.tableWidget.setColumnCount(len(headers))
        self.tableWidget.setHorizontalHeaderLabels(headers)
        self.tableWidget.setRowCount(len(data))

        for row_idx, row_data in enumerate(data):
            for col_idx, cell_data in enumerate(row_data):
                item = QtWidgets.QTableWidgetItem(str(cell_data) if cell_data is not None else "")
                self.tableWidget.setItem(row_idx, col_idx, item)

        self.tableWidget.resizeColumnsToContents()

    def search(self):
            search_text = self.lineEdit_ara.text().strip().lower()
            if not search_text:
                QtWidgets.QMessageBox.information(self, "Bilgi", "Lütfen arama metni giriniz.")
                return

            headers, data = self.load_excel()
            if not data:
                return

            filtered = []
            for row in data:
                # Her hücreyi kontrol et, ancak 'startswith' ile
                if any(cell and str(cell).lower().startswith(search_text) for cell in row):
                    filtered.append(row)

            if not filtered:
                QtWidgets.QMessageBox.information(self, "Sonuç", "Arama kriterlerine uygun kayıt bulunamadı.")
                self.tableWidget.clear()
                self.tableWidget.setRowCount(0)
                self.tableWidget.setColumnCount(0)
                return

            self.display_data(headers, filtered)

    def interviews(self):
        headers, data = self.load_excel()
        if not data:
            return
        self.display_data(headers, data)

    def multiple_tabs_changed(self, index):
        if self.initializing:
            # Açılışta otomatik tetiklenmeyi engelle
            return

        secilen = self.comboBox_cokluSekme.currentText()
        headers, data = self.load_excel()
        if not data:
            self.tableWidget.clear()
            self.tableWidget.setRowCount(0)
            self.tableWidget.setColumnCount(0)
            return

        if secilen == "Hepsi":
            self.display_data(headers, data)
            return

        filtre_sutun_adi = "VIT projesinin tamamına katılması uygun olur"
        try:
            filtre_idx = headers.index(filtre_sutun_adi)
        except ValueError:
            self.display_data(headers, data)
            return

        filtered = [row for row in data if row[filtre_idx] == secilen]

        if not filtered:
            QtWidgets.QMessageBox.information(self, "Sonuç", f"'{secilen}' için kayıt bulunamadı.")
            self.tableWidget.clear()
            self.tableWidget.setRowCount(0)
            self.tableWidget.setColumnCount(0)
            return

        self.display_data(headers, filtered)

    def fill_combobox_categories(self):
        headers, data = self.load_excel()
        if not data:
            self.comboBox_cokluSekme.clear()
            self.comboBox_cokluSekme.addItem("Hepsi")
            return

        filtre_sutun_adi = "VIT projesinin tamamına katılması uygun olur"
        try:
            filtre_idx = headers.index(filtre_sutun_adi)
        except ValueError:
            self.comboBox_cokluSekme.clear()
            self.comboBox_cokluSekme.addItem("Hepsi")
            return

        unique_values = sorted(set(row[filtre_idx] for row in data if row[filtre_idx] is not None))

        self.comboBox_cokluSekme.clear()
        self.comboBox_cokluSekme.addItem("Hepsi")
        for val in unique_values:
            self.comboBox_cokluSekme.addItem(str(val))

    def go_back(self):
        try:
            if self.sayfa:
                self.sayfa.show()
        except Exception as e:
            QtWidgets.QMessageBox.warning(self, "Hata", f"Geri dönülecek sayfa açılamadı: {str(e)}")
        self.close()

    def exitp(self):
        QtWidgets.QApplication.quit()
    def diger_sayfa(self,page):
        self.sayfa = page
