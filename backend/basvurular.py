from PyQt6.QtWidgets import QMainWindow, QTableWidgetItem, QMessageBox # QMessageBox'u ekledik
from PyQt6 import uic, QtWidgets
import os 
import openpyxl
from backend.find_excel import get_excel_files_from_latest_vit # find_excel.py import'ı yerinde
import sys


class BasvurularPage(QMainWindow):
    
    def __init__(self):
        super().__init__()

        # UI ve Excel dosyaları için TEK VE TUTARLI BİR app_root_path tanımlaması yapıyoruz
        if getattr(sys, 'frozen', False):
            # PyInstaller ile paketlenmişse, _MEIPASS geçici dizinin yolunu verir
            self.app_root_path = sys._MEIPASS
        else:
            # Normal Python betiği olarak çalışıyorsa
            current_script_dir = os.path.dirname(os.path.abspath(__file__))
            self.app_root_path = os.path.abspath(os.path.join(current_script_dir, os.pardir))
            # print(f"Normal çalışma app_root_path: {self.app_root_path}") # Debug için

        # UI dosyasının tam yolunu oluştur
        ui_file_name = "basvurular_page_python.ui"
        ui_file_path = os.path.join(self.app_root_path, "ui", ui_file_name) # self.app_root_path kullanıldı

        # UI dosyasını yükle
        try:
            uic.loadUi(ui_file_path, self)
        except Exception as e:
            QMessageBox.critical(self, "UI Yükleme Hatası", f"UI dosyası yüklenemedi: {ui_file_path}\nHata: {e}")
            sys.exit(1) # Hata durumunda uygulamayı kapat

        self.sayfa = None
        self.kullanici = ""
        
        # Excel tablo işlemleri (temel düzey) - Tek seferlik yükleme
        try:
            # get_excel_files_from_latest_vit fonksiyonuna doğru kök yolu iletin
            excel_files = get_excel_files_from_latest_vit(self.app_root_path) # self.app_root_path kullanıldı
            basvurular_path = excel_files.get("Basvurular")

            if not basvurular_path or not os.path.exists(basvurular_path):
                raise FileNotFoundError(f"Basvurular.xlsx dosyası bulunamadı. Lütfen dosyanın 'data' klasöründe olduğundan emin olun. Yol: {basvurular_path}")

            self.excel_path = basvurular_path # Excel dosya yolunu kaydediyoruz
            dosya_basvurular = openpyxl.load_workbook(self.excel_path)
            self.sayfa_basvurular = dosya_basvurular.active # sayfa_basvurular'ı doğru şekilde set et
            self.satir_sayisi = self.sayfa_basvurular.max_row
            self.sutun_sayisi = self.sayfa_basvurular.max_column

        except Exception as e:
            QtWidgets.QMessageBox.critical(self, "Hata", f"Excel dosyası yüklenemedi: {e}")
            self.excel_path = None # Hata durumunda path'i None yap
            self.sayfa_basvurular = None # Sayfayı da None yap
            # return # Hata durumunda uygulamanızın kapanmasını istiyorsanız burayı açın


        # ui öğelerine erişim (UI'daki isimlendirmelerinize göre kontrol edin!)
        self.pushButton_tumBasvurular = self.findChild(QtWidgets.QPushButton,"pushButton_tumBasvurular")
        self.pushButton_tanimlananMentorGorusmesi = self.findChild(QtWidgets.QPushButton,"pushButton_tanimlananMentorGorusmesi")
        self.pushButton_tanimlanmayanMentorGorusmesi = self.findChild(QtWidgets.QPushButton,"pushButton_tanimlanmayanMentorGorusmesi")
        self.pushButton_oncekiVitKontrol = self.findChild(QtWidgets.QPushButton,"pushButton_oncekiVitKontrol")
        self.pushButton_basvuruFiltrele = self.findChild(QtWidgets.QPushButton,"pushButton_basvuruFiltrele")
        self.pushButton_mukerrerKayit = self.findChild(QtWidgets.QPushButton,"pushButton_mukerrerKayit")
        self.pushButton_farkliKayit = self.findChild(QtWidgets.QPushButton,"pushButton_farkliKayit")
        self.pushButton_tercihler = self.findChild(QtWidgets.QPushButton,"pushButton_tercihler")
        self.pushButton_exit = self.findChild(QtWidgets.QPushButton,"pushButton_exit")
        self.pushButton_ara = self.findChild(QtWidgets.QPushButton,"pushButton_ara")
        self.lineEdit_ara = self.findChild(QtWidgets.QLineEdit,"lineEdit_ara")
        self.tableWidget_goster = self.findChild(QtWidgets.QTableWidget,"tableWidget_goster")

        # TableWidget'ı başlangıçta temizle (eğer UI'da zaten tanımlıysa)
        if self.tableWidget_goster:
            self.tableWidget_goster.setColumnCount(0)
            self.tableWidget_goster.setRowCount(0)

        # olaylar
        self.pushButton_tumBasvurular.clicked.connect(self.tumbasvulari_getir)
        self.pushButton_tanimlananMentorGorusmesi.clicked.connect(self.mentor_gosrusmesi_tanimlananlar)
        self.pushButton_tanimlanmayanMentorGorusmesi.clicked.connect(self.mentor_gosrusmesi_tanimlanmayanlar)
        self.pushButton_oncekiVitKontrol.clicked.connect(self.onceki_vit_kontrol)
        self.pushButton_basvuruFiltrele.clicked.connect(self.basvurulari_filtrele)
        self.pushButton_mukerrerKayit.clicked.connect(self.mukerrer_kayitlar)
        self.pushButton_farkliKayit.clicked.connect(self.farkli_kayitlar)
        self.pushButton_tercihler.clicked.connect(self.tercihler_sayfasina_don)
        self.pushButton_exit.clicked.connect(self.cikis_yap)
        self.pushButton_ara.clicked.connect(self.arama_yapilamasi)
        
    # --- Yardımcı Fonksiyonlar ---
    def load_excel_data(self):
        """
        Excel dosyasını yükler ve başlıklar ile verileri döndürür.
        """
        if not self.excel_path or not os.path.exists(self.excel_path):
            QMessageBox.warning(self, "Hata", "Basvurular.xlsx dosyası bulunamadı veya yolu geçersiz.")
            return [], []

        try:
            wb = openpyxl.load_workbook(self.excel_path)
            sheet = wb.active
            headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
            data = []
            for row in sheet.iter_rows(min_row=2, values_only=True):
                data.append(list(row))
            return headers, data
        except Exception as e:
            QMessageBox.critical(self, "Hata", f"Excel dosyası okunurken hata oluştu: {e}")
            return [], []

    def display_data(self, headers, data):
        """
        Verilen başlıklar ve verilerle tableWidget_goster'ı doldurur.
        """
        self.tableWidget_goster.clearContents()
        self.tableWidget_goster.setColumnCount(len(headers))
        self.tableWidget_goster.setHorizontalHeaderLabels(headers)
        self.tableWidget_goster.setRowCount(len(data))

        for row_idx, row_data in enumerate(data):
            for col_idx, cell_data in enumerate(row_data):
                item = QTableWidgetItem(str(cell_data) if cell_data is not None else "")
                self.tableWidget_goster.setItem(row_idx, col_idx, item)
        self.tableWidget_goster.resizeColumnsToContents()

    # --- Sinyal Fonksiyonları ---

    def diger_sayfa(self, page):
        self.sayfa = page

    def tercihler_sayfasina_don(self):
        if self.sayfa:
            self.sayfa.show()
        self.hide()

    def tumbasvulari_getir(self):
        headers, data = self.load_excel_data() # Yeni yardımcı fonksiyonu kullan
        if headers and data:
            self.display_data(headers, data)

    def mentor_gosrusmesi_tanimlananlar(self):
        headers, data = self.load_excel_data()
        if not data:
            return

        mentor_gorusmesi_index = -1
        # Tüm başlıkları küçük harfe çevirip boşlukları temizleyerek arama yapıyoruz
        for i, header in enumerate(headers):
            if header and header.strip().lower() == "mentor gorusmesi": # Küçük harfe çevirip boşlukları temizleyip karşılaştır
                mentor_gorusmesi_index = i
                break
        
        if mentor_gorusmesi_index == -1:
            QMessageBox.warning(self, "Hata", "'Mentor gorusmesi' sütunu bulunamadı. Lütfen Excel'deki sütun adını kontrol edin.")
            return

        filtered_data = [row for row in data if row[mentor_gorusmesi_index] and str(row[mentor_gorusmesi_index]).strip().upper() == "OK"]
        
        if not filtered_data:
            QMessageBox.information(self, "Bilgi", "Tanımlanmış mentor görüşmesi olan aday bulunamadı.")
            self.tableWidget_goster.clearContents()
            self.tableWidget_goster.setRowCount(0)
            self.tableWidget_goster.setColumnCount(0)
            return

        self.display_data(headers, filtered_data)

    def mentor_gosrusmesi_tanimlanmayanlar(self):
            print("\n--- mentor_gosrusmesi_tanimlanmayanlar fonksiyonu çağrıldı ---") # Debug: Fonksiyonun çağrıldığını onayla

            headers, data = self.load_excel_data()
            if not data:
                print("Debug: Veri yüklenemedi veya boş.") # Debug
                return

            print(f"Debug: Yüklenen başlıklar: {headers}") # Debug: Başlıkları görelim

            mentor_gorusmesi_index = -1
            # Tüm başlıkları küçük harfe çevirip boşlukları temizleyerek arama yapıyoruz
            for i, header in enumerate(headers):
                normalized_header = header.strip().lower() if header else ""
                print(f"Debug: Başlık {i}: '{header}' -> Normalleştirilmiş: '{normalized_header}'") # Debug: Her başlığı ve normalleştirilmiş halini gör
                if normalized_header == "mentor gorusmesi": # Küçük harfe çevirip boşlukları temizleyip karşılaştır
                    mentor_gorusmesi_index = i
                    print(f"Debug: 'Mentor gorusmesi' sütunu bulundu, indeks: {mentor_gorusmesi_index}") # Debug
                    break

            if mentor_gorusmesi_index == -1:
                print("Debug: 'Mentor gorusmesi' sütunu bulunamadı.") # Debug
                QMessageBox.warning(self, "Hata", "'Mentor gorusmesi' sütunu bulunamadı. Lütfen Excel'deki sütun adını kontrol edin.")
                return

            filtered_data = [row for row in data if row[mentor_gorusmesi_index] and str(row[mentor_gorusmesi_index]).strip().upper() == "ATANMADI"]
            
            print(f"Debug: Toplam veri satırı: {len(data)}") # Debug
            print(f"Debug: Filtrelenmiş veri satırı: {len(filtered_data)}") # Debug

            if not filtered_data:
                QMessageBox.information(self, "Bilgi", "Tanımlanmamış mentor görüşmesi olan aday bulunamadı.")
                self.tableWidget_goster.clearContents()
                self.tableWidget_goster.setRowCount(0)
                self.tableWidget_goster.setColumnCount(0)
                return

            self.display_data(headers, filtered_data)
            print("Debug: Veriler tableWidget'a yüklendi.") # Debug

    def onceki_vit_kontrol(self):
        import glob
        from collections import defaultdict

        adaylar = defaultdict(lambda: {"veri": None, "vitler": set(), "headers": []}) # headers'ı da saklayalım

        # app_root_path'i kullanarak 'data' dizinine doğru yolu oluştur
        data_dizini = os.path.join(self.app_root_path, "data")
        vit_klasorleri = [f for f in glob.glob(os.path.join(data_dizini, "VIT*")) if os.path.isdir(f)]

        if not vit_klasorleri:
            QMessageBox.information(self, "Bilgi", "Hiç VIT klasörü bulunamadı.")
            return

        for vit_klasoru in vit_klasorleri:
            vit_adi = os.path.basename(vit_klasoru)
            excel_path = os.path.join(vit_klasoru, "Basvurular.xlsx")
            if not os.path.exists(excel_path):
                continue

            try:
                kitap = openpyxl.load_workbook(excel_path)
                sayfa = kitap.active

                current_headers = [str(sayfa.cell(1, col).value).strip() if sayfa.cell(1, col).value is not None else "" for col in range(1, sayfa.max_column + 1)]
                
                # Sadece ilk okuduğumuz VIT'in başlıklarını genel başlıklar olarak alalım
                if not adaylar[list(adaylar.keys())[0] if adaylar else ""][ "headers"]:
                    adaylar[list(adaylar.keys())[0] if adaylar else ""]["headers"] = current_headers


                eposta_basliklari = ["e-posta", "email", "mail", "mail adresiniz"]
                email_index = None
                for i, baslik in enumerate(current_headers): # current_headers'ı kullan
                    if baslik and baslik.lower() in eposta_basliklari:
                        email_index = i # Python index (0-indexed)
                        break

                if email_index is None: # email_index'i burada kontrol edin
                    continue

                for row_idx in range(2, sayfa.max_row + 1):
                    email = sayfa.cell(row_idx, email_index + 1).value # Excel index (1-indexed)
                    if email:
                        email = str(email).strip().lower()
                        satir_verisi = [sayfa.cell(row_idx, col).value for col in range(1, sayfa.max_column + 1)]

                        if not adaylar[email]["veri"]:
                            adaylar[email]["veri"] = satir_verisi
                            adaylar[email]["headers"] = current_headers # Bu adayın verisinin başlıklarını da kaydet
                        adaylar[email]["vitler"].add(vit_adi)

            except Exception as e:
                QMessageBox.warning(self, "Uyarı", f"'{vit_adi}' klasöründeki Basvurular.xlsx dosyası okunurken hata oluştu: {e}")

        ortak_adaylar = {email: bilgi for email, bilgi in adaylar.items() if len(bilgi["vitler"]) > 1}

        if ortak_adaylar:
            # Dinamik başlıkları al (ilk adayın veri başlıklarını kullanabiliriz)
            sample_email = next(iter(ortak_adaylar))
            headers_to_display = ortak_adaylar[sample_email]["headers"] + ["Başvurduğu VIT'ler"]
            
            display_data_list = []
            for email, bilgi in ortak_adaylar.items():
                row_to_display = list(bilgi["veri"]) # Orjinal verinin bir kopyası
                row_to_display.append(", ".join(sorted(list(bilgi["vitler"]))))
                display_data_list.append(row_to_display)
            
            self.display_data(headers_to_display, display_data_list)
        else:
            QMessageBox.information(self, "Bilgi", "Hiçbir aday birden fazla VIT'e başvurmamış.")

    def mukerrer_kayitlar(self):
        import glob
        from collections import defaultdict

        # app_root_path'i kullanarak 'data' dizinine doğru yolu oluştur
        data_dizini = os.path.join(self.app_root_path, "data")
        vit_klasorleri = sorted(
            [f for f in glob.glob(os.path.join(data_dizini, "VIT*")) if os.path.isdir(f)],
            reverse=True
        )

        if not vit_klasorleri:
            QMessageBox.information(self, "Bilgi", "Hiç VIT klasörü bulunamadı.")
            return

        # Sadece en son VIT'in Basvurular.xlsx dosyasını kullan
        son_vit_klasoru = vit_klasorleri[0]
        basvurular_path = os.path.join(son_vit_klasoru, "Basvurular.xlsx")

        if not os.path.exists(basvurular_path):
            QMessageBox.information(self, "Bilgi", f"{son_vit_klasoru} içinde Basvurular.xlsx dosyası bulunamadı.")
            return

        try:
            kitap = openpyxl.load_workbook(basvurular_path)
            sayfa = kitap.active

            basliklar = [str(sayfa.cell(1, col).value).strip() if sayfa.cell(1, col).value is not None else "" for col in range(1, sayfa.max_column + 1)]
            
            eposta_basliklari = ["e-posta", "email", "mail", "mail adresiniz"]
            email_index = None
            for i, baslik in enumerate(basliklar):
                if baslik and baslik.lower() in eposta_basliklari:
                    email_index = i + 1
                    break

            if not email_index:
                QMessageBox.information(self, "Bilgi", "E-posta sütunu bulunamadı.")
                return

            adaylar = defaultdict(list)
            for row_idx in range(2, sayfa.max_row + 1):
                email_cell_value = sayfa.cell(row_idx, email_index).value
                if email_cell_value:
                    email = str(email_cell_value).strip().lower()
                    satir_verisi = [sayfa.cell(row_idx, col).value for col in range(1, sayfa.max_column + 1)]
                    adaylar[email].append(satir_verisi)

            mukerrer_adaylar = {email: kayitlar for email, kayitlar in adaylar.items() if len(kayitlar) > 1}

            if not mukerrer_adaylar:
                QMessageBox.information(self, "Bilgi", "Mükerrer kayıt bulunamadı.")
                self.tableWidget_goster.clearContents()
                self.tableWidget_goster.setRowCount(0)
                self.tableWidget_goster.setColumnCount(0)
                return
            
            display_data_list = []
            for email, kayitlar in mukerrer_adaylar.items():
                display_data_list.extend(kayitlar)

            self.display_data(basliklar, display_data_list)

        except Exception as e:
            QMessageBox.critical(self, "Hata", f"Mükerrer kayıtlar yüklenirken hata oluştu: {e}")

    def farkli_kayitlar(self):
        import glob
        from collections import defaultdict

        adaylar = defaultdict(lambda: {"veri": None, "vitler": set(), "headers": []})

        # app_root_path'i kullanarak 'data' dizinine doğru yolu oluştur
        data_dizini = os.path.join(self.app_root_path, "data")
        vit_klasorleri = [f for f in glob.glob(os.path.join(data_dizini, "VIT*")) if os.path.isdir(f)]

        if not vit_klasorleri:
            QMessageBox.information(self, "Bilgi", "Hiç VIT klasörü bulunamadı.")
            return

        basliklar_genel = [] # Tüm VIT'lerden başlıkları almak için
        
        for vit_klasoru in vit_klasorleri:
            vit_adi = os.path.basename(vit_klasoru)
            excel_path = os.path.join(vit_klasoru, "Basvurular.xlsx")
            if not os.path.exists(excel_path):
                continue

            try:
                kitap = openpyxl.load_workbook(excel_path)
                sayfa = kitap.active

                current_headers = [str(sayfa.cell(1, col).value).strip() if sayfa.cell(1, col).value is not None else "" for col in range(1, sayfa.max_column + 1)]
                
                # İlk okuduğumuz VIT'in başlıklarını genel başlıklar olarak alalım (veya daha gelişmiş bir başlık birleştirme yapabilirsiniz)
                if not basliklar_genel:
                    basliklar_genel = current_headers

                eposta_basliklari = ["e-posta", "email", "mail", "mail adresiniz"]
                email_index = None
                for i, baslik in enumerate(current_headers):
                    if baslik and baslik.lower() in eposta_basliklari:
                        email_index = i 
                        break

                if email_index is None:
                    continue

                for row_idx in range(2, sayfa.max_row + 1):
                    email = sayfa.cell(row_idx, email_index + 1).value
                    if email:
                        email = str(email).strip().lower()
                        satir_verisi = [sayfa.cell(row_idx, col).value for col in range(1, sayfa.max_column + 1)]

                        if not adaylar[email]["veri"]:
                            adaylar[email]["veri"] = satir_verisi
                            adaylar[email]["headers"] = current_headers # Bu adayın verisinin başlıklarını da kaydet
                        adaylar[email]["vitler"].add(vit_adi)

            except Exception as e:
                QMessageBox.warning(self, "Uyarı", f"'{vit_adi}' klasöründeki Basvurular.xlsx dosyası okunurken hata oluştu: {e}")

        farkli_adaylar = {email: bilgi for email, bilgi in adaylar.items() if len(bilgi["vitler"]) == 1}

        if farkli_adaylar:
            # Dinamik başlıkları al (ilk adayın veri başlıklarını kullanabiliriz)
            sample_email = next(iter(farkli_adaylar))
            headers_to_display = farkli_adaylar[sample_email]["headers"] + ["Başvurduğu VIT"]
            
            display_data_list = []
            for email, bilgi in farkli_adaylar.items():
                row_to_display = list(bilgi["veri"])
                row_to_display.append(list(bilgi["vitler"])[0]) # Tek olduğu için ilk elemanı al
                display_data_list.append(row_to_display)
            
            self.display_data(headers_to_display, display_data_list)
        else:
            QMessageBox.information(self, "Bilgi", "Tüm adaylar birden fazla VIT'e başvurmuş.")

    def basvurulari_filtrele(self):
        headers, data = self.load_excel_data()
        if not data:
            return

        eposta_basliklari = ["e-posta", "email", "mail", "mail adresiniz"]
        email_index = None
        for i, baslik in enumerate(headers):
            if baslik and baslik.lower() in eposta_basliklari:
                email_index = i
                break

        if email_index is None:
            QMessageBox.information(self, "Bilgi", "E-posta sütunu bulunamadı, filtreleme yapılamıyor.")
            return

        benzersiz_kayitlar = {}
        for row_data in data:
            email = row_data[email_index]
            if email:
                email = str(email).strip().lower()
                if email not in benzersiz_kayitlar:
                    benzersiz_kayitlar[email] = row_data
        
        display_data_list = list(benzersiz_kayitlar.values())

        if not display_data_list:
            QMessageBox.information(self, "Bilgi", "Filtrelenecek benzersiz kayıt bulunamadı.")
            self.tableWidget_goster.clearContents()
            self.tableWidget_goster.setRowCount(0)
            self.tableWidget_goster.setColumnCount(0)
            return

        self.display_data(headers, display_data_list)

    def cikis_yap(self):
        self.close() 
    
    def arama_yapilamasi(self):
        aranan_kelime = self.lineEdit_ara.text().strip().lower()
        if not aranan_kelime:
            QMessageBox.information(self, "Bilgi", "Lütfen arama yapmak için bir isim giriniz.")
            return

        headers, data = self.load_excel_data()
        if not data:
            return

        # Başlık satırını al
        basliklar_lower = [h.lower() for h in headers]

        # İsim ve soyisim kolonlarının indekslerini bul
        ad_index = None
        soyad_index = None
        for i, baslik in enumerate(basliklar_lower):
            if "ad" in baslik and "soy" not in baslik:
                ad_index = i
            elif "soyad" in baslik:
                soyad_index = i

        if ad_index is None or soyad_index is None:
            QMessageBox.critical(self, "Hata", "Ad veya Soyad sütunları bulunamadı.")
            return

        eslesen_satirlar = []
        for row_data in data:
            ad = str(row_data[ad_index] or "").strip().lower()
            soyad = str(row_data[soyad_index] or "").strip().lower()
            
            if ad.startswith(aranan_kelime) or soyad.startswith(aranan_kelime):
                eslesen_satirlar.append(row_data)

        if not eslesen_satirlar:
            QMessageBox.information(self, "Bilgi", "Eşleşen kayıt bulunamadı.")
            self.tableWidget_goster.clearContents()
            self.tableWidget_goster.setRowCount(0)
            self.tableWidget_goster.setColumnCount(0)
            return

        self.display_data(headers, eslesen_satirlar)